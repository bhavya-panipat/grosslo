"""
tests/test_export_basis_flags.py — D-S3: the basis flags must reach the payment
instruction, not only the Finance screen.

`review_queue._row_to_dict()` attaches `tax_basis_flag` and
`treasury_basis_flag` to every row it returns. The per-row export route reads its
row through `get_submission()`, so it HOLDS both and used neither — `app.py`
contained no reference to `tax_basis` at all.

Every reason string here is asserted against `review_queue`'s own constant, never
a copied literal. A copied string would pass while the route emitted stale text,
which is the same class of defect as a figure copied out of a run.

See EXPORT_BASIS_FLAG_DESIGN.md. Note §2: the exportable-and-flagged population
is measured at zero and the flagged set is closed, so this is defense-in-depth
rather than a repair. These tests construct the flagged state deliberately —
which is also the only way it can now occur.
"""

import io
import unittest

import review_queue
from openpyxl import load_workbook

from tests.test_review_workflow import ReviewQueueTestCase, _client, _login_as


ROW = {
    "employee_name": "Basis Flag", "ctc": 1_800_000, "rent_paid": 0, "city": "metro",
    "nps_opted": True, "band_min": 1_700_000, "band_max": 1_900_000,
    "work_location": "karnataka",
    "bank_account_number": "1234567890", "ifsc": "HDFC0000001",
    "email": "basis@test.com",
}

CORRECTION_ROW = {
    "employee_name": "Correction Flag", "ctc": 4_000_000, "rent_paid": 500_000,
    "city": "metro", "nps_opted": True,
    "current_structure": {
        "basic": 1_800_000, "hra": 900_000, "lta": 100_000,
        "special_allowance": 300_000, "employer_pf": 900_000, "employer_nps": 280_000,
    },
}


class _ExportBase(ReviewQueueTestCase):

    def setUp(self):
        super().setUp()
        self.client = _client()
        _login_as(self.client, "finance")

    def _submit_and_approve(self, row):
        resp = self.client.post("/api/submissions", json={"source": "single", "row": row})
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        sid = resp.get_json()["submission_id"]
        decide = self.client.post(f"/api/submissions/{sid}/rows/0/decide",
                                  json={"decision": "approve"})
        self.assertEqual(decide.status_code, 200)
        return sid

    def _force_pre_fix_basis(self, submission_id):
        """
        The only way a flagged row can now exist: tax_basis is written at insert
        as the one basis treated as correct, so the flagged set is closed
        (EXPORT_BASIS_FLAG_DESIGN.md §2.2). NULL is what a row predating the
        column reads as.
        """
        tenant_id = review_queue._checked_tenant_id(1)
        with review_queue._conn(tenant_id) as conn:
            conn.execute(
                "UPDATE submission_rows SET tax_basis = NULL "
                "WHERE tenant_id = %s AND submission_id = %s AND row_index = 0",
                (tenant_id, submission_id))

    def _assert_row_is_flagged(self, submission_id):
        """Precondition, not an assertion about the route: if the fixture stops
        producing a flagged row, the tests below would pass vacuously."""
        tenant_id = review_queue._checked_tenant_id(1)
        row = review_queue.get_submission(tenant_id, submission_id)["rows"][0]
        self.assertIsNotNone(row["tax_basis_flag"],
                             "fixture no longer produces a flagged row")
        return row


class TestTheFlagReachesThePayoutPayload(_ExportBase):

    def test_a_flagged_row_carries_the_flag_and_the_warning(self):
        sid = self._submit_and_approve(dict(ROW))
        self._force_pre_fix_basis(sid)
        self._assert_row_is_flagged(sid)

        body = self.client.post(f"/api/submissions/{sid}/rows/0/export").get_json()
        self.assertIsNotNone(body.get("tax_basis_flag"))
        self.assertEqual(body["tax_basis_flag"]["reason"],
                         review_queue.PRE_NPS_FIX_REASON)
        # Two surfaces, as WARNING_DO_NOT_UPLOAD already gets in this route: a
        # caveat on a payment instruction has to survive being piped or saved.
        self.assertIn(review_queue.PRE_NPS_FIX_REASON,
                      body["WARNING_BASIS_SUPERSEDED"])

    def test_an_unflagged_row_carries_neither_key(self):
        # The both-states pair. Without this, always emitting would pass.
        sid = self._submit_and_approve(dict(ROW))
        body = self.client.post(f"/api/submissions/{sid}/rows/0/export").get_json()
        self.assertNotIn("tax_basis_flag", body)
        self.assertNotIn("treasury_basis_flag", body)
        self.assertNotIn("WARNING_BASIS_SUPERSEDED", body)

    def test_the_figures_are_identical_whether_or_not_the_row_is_flagged(self):
        # D1-4 and D-T3: nothing is recomputed and nothing stored is rewritten.
        # The flag is a caveat on the figures, not a change to them.
        clean_sid = self._submit_and_approve(dict(ROW, employee_name="Clean",
                                                  email="clean@test.com"))
        clean = self.client.post(
            f"/api/submissions/{clean_sid}/rows/0/export").get_json()

        flagged_sid = self._submit_and_approve(dict(ROW, employee_name="Flagged",
                                                   email="flagged@test.com"))
        self._force_pre_fix_basis(flagged_sid)
        flagged = self.client.post(
            f"/api/submissions/{flagged_sid}/rows/0/export").get_json()

        self.assertEqual(flagged["payouts"][0]["amount"],
                         clean["payouts"][0]["amount"])
        self.assertEqual(flagged["payout_basis"], clean["payout_basis"])
        self.assertEqual(flagged["treasury_forecast"], clean["treasury_forecast"])

    def test_the_route_still_exports_rather_than_refusing(self):
        # D1-7(b): the reasons say so, the route does not block. A refusal would
        # be a new gate on work a human has already approved.
        sid = self._submit_and_approve(dict(ROW))
        self._force_pre_fix_basis(sid)
        resp = self.client.post(f"/api/submissions/{sid}/rows/0/export")
        self.assertEqual(resp.status_code, 200)


class TestTheFlagReachesTheWorkbookExport(_ExportBase):
    """
    The correction path returns an XLSX, so the flag cannot ride in a JSON body.
    EXPORT_BASIS_FLAG_DESIGN.md §7 left this open; it is decided here as the
    workbook itself plus a header, because a saved file outlives its response
    headers — the same reasoning that put the JSON caveat in the body.
    """

    def test_a_flagged_correction_export_says_so_inside_the_file(self):
        sid = self._submit_and_approve(dict(CORRECTION_ROW))
        self._force_pre_fix_basis(sid)
        self._assert_row_is_flagged(sid)

        resp = self.client.post(f"/api/submissions/{sid}/rows/0/export")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheet", resp.content_type)

        wb = load_workbook(io.BytesIO(resp.data))
        readme = " ".join(str(c.value) for row in wb["Read Me"].iter_rows()
                          for c in row if c.value is not None)
        self.assertIn(review_queue.PRE_NPS_FIX_REASON, readme)

    def test_a_flagged_correction_export_also_sets_a_header(self):
        sid = self._submit_and_approve(dict(CORRECTION_ROW))
        self._force_pre_fix_basis(sid)
        resp = self.client.post(f"/api/submissions/{sid}/rows/0/export")
        self.assertIn("X-Basis-Superseded", resp.headers)

    def test_an_unflagged_correction_export_is_unchanged(self):
        # The both-states pair for the workbook path, and it also protects the
        # existing Read Me content from gaining an empty warning row.
        sid = self._submit_and_approve(dict(CORRECTION_ROW,
                                            employee_name="Clean Correction"))
        resp = self.client.post(f"/api/submissions/{sid}/rows/0/export")
        self.assertEqual(resp.status_code, 200)
        self.assertNotIn("X-Basis-Superseded", resp.headers)
        wb = load_workbook(io.BytesIO(resp.data))
        readme = " ".join(str(c.value) for row in wb["Read Me"].iter_rows()
                          for c in row if c.value is not None)
        self.assertNotIn(review_queue.PRE_NPS_FIX_REASON, readme)
        self.assertNotIn("superseded", readme.lower())

    def test_the_existing_honesty_label_survives_alongside_the_warning(self):
        # Two caveats must coexist: the template label is about the FILE SHAPE,
        # the basis warning is about the FIGURES. Neither replaces the other.
        from salary_revision_export import TEMPLATE_HONESTY_LABEL
        sid = self._submit_and_approve(dict(CORRECTION_ROW))
        self._force_pre_fix_basis(sid)
        resp = self.client.post(f"/api/submissions/{sid}/rows/0/export")
        wb = load_workbook(io.BytesIO(resp.data))
        readme = " ".join(str(c.value) for row in wb["Read Me"].iter_rows()
                          for c in row if c.value is not None)
        self.assertIn(TEMPLATE_HONESTY_LABEL, readme)
        self.assertIn(review_queue.PRE_NPS_FIX_REASON, readme)
        self.assertIn("X-Template-Honesty-Label", resp.headers)


if __name__ == "__main__":
    unittest.main()
